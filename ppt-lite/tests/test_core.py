"""集成测试：端到端（生成 pptx → 上传 → 解析 → finalize → 检索/审核）
+ 故障注入矩阵关键项（lite-solution §9）。"""
from __future__ import annotations

import io
import json
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from ppt_lite.config import Config
from ppt_lite.db import Database, Conflict, InvariantViolation
from ppt_lite.precheck import PrecheckError, precheck
from ppt_lite.store import Store
from ppt_lite.worker import Worker


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("PPT_LITE_DATA", str(tmp_path / "data"))
    cfg = Config()
    cfg.ensure_dirs()
    db = Database(cfg)
    store = Store(db, cfg)
    return cfg, db, store


def make_pptx(path: Path, n_slides: int = 3) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(n_slides):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        tf = tb.text_frame
        tf.text = f"第{i+1}页 战略规划标题"
        p = tf.add_paragraph()
        p.text = f"内容文本 第{i+1}页 详细说明"
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size = Pt(18)
    prs.save(str(path))
    return path.read_bytes()


def upload(cfg, store, data: bytes, name="t.pptx", category="sample") -> int:
    """模拟 API 上传路径（预检→暂存→register_upload）。"""
    import hashlib
    from ppt_lite.fslock import fslock
    import uuid
    res = precheck(data, soffice_available=False)
    assert res.kind == "pptx"
    staging = cfg.tmp_dir / f"upload-{uuid.uuid4().hex}"
    staging.mkdir(parents=True, exist_ok=True)
    raw_src = staging / "blob"
    raw_src.write_bytes(data)
    sha = hashlib.sha256(data).hexdigest()
    raw_rel = f"raw/{sha}.pptx"
    fid, _ = store.register_upload(sha, name, category, raw_rel, None, raw_src, None)
    from ppt_lite.fslock import to_trash
    to_trash(staging, cfg.trash_dir)
    return fid


def drain(worker: Worker):
    worker._wake.set()
    worker._run_once = True
    # 手动循环跑完队列（测试不启动线程）
    while True:
        job = worker.store.claim()
        if job is None:
            return
        worker._process(*job)


# ------------------------------------------------------------------ 端到端
def test_end_to_end(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "src.pptx", n_slides=3)
    fid = upload(cfg, store, data)

    f = db.one("select * from file where id=?", fid)
    assert f["status"] == "queued"
    ex = db.one("select * from extraction where file_id=?", fid)
    assert ex["status"] == "queued"

    w = Worker(db, cfg)
    drain(w)

    ex = db.one("select * from extraction where file_id=?", fid)
    assert ex["status"] == "done", ex["error"]
    f = db.one("select * from file where id=?", fid)
    assert f["status"] == "done"
    slides = db.all("select * from v_active_slide where file_id=? order by no", fid)
    assert len(slides) == 3
    assert "战略规划标题" in slides[0]["search_text"]
    # 预览目录（LO 缺失 → render_error 降级：png 路径已定，文件可能缺失）
    assert slides[0]["png"] == f"previews/{ex['id']}/slides/1.png"

    # 检索：≥3 连续中文走 FTS；短查询走 LIKE
    from ppt_lite.search import search_slides
    hits = search_slides(db, "战略规划标题")
    assert len(hits) == 3
    hits = search_slides(db, "规划")     # 2 字 → LIKE
    assert len(hits) == 3
    hits = search_slides(db, "不存在的词组xyz")
    assert len(hits) == 0

    # 审核 → forbidden 过滤
    sid = slides[0]["id"]
    with db.tx() as cur:
        cur.execute("insert into slide_review(file_id,no,status) values(?,?, 'forbidden')", (fid, 1))
    hits = search_slides(db, "战略规划标题")
    assert len(hits) == 2


# ------------------------------------------------------------------ 矩阵：重跑保留人工成果
def test_reprocess_keeps_review(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "src.pptx", n_slides=2)
    fid = upload(cfg, store, data)
    w = Worker(db, cfg)
    drain(w)
    with db.tx() as cur:
        cur.execute("insert into slide_review(file_id,no,status,slide_type) values(?,1,'approved','cover')", (fid,))

    # 重跑（新 extraction）
    with db.tx() as cur:
        cur.execute("insert into extraction(file_id,status) values(?, 'queued')", (fid,))
    drain(w)

    exs = db.all("select * from extraction where file_id=?", fid)
    dones = [e for e in exs if e["status"] == "done"]
    assert len(dones) == 1                      # 生效 done 恰好一个（首个若失败则 failed 行保留属预期）
    sr = db.one("select * from slide_review where file_id=? and no=1", fid)
    assert sr["status"] == "approved" and sr["slide_type"] == "cover"   # 人工成果保留
    assert db.one("select count(*) c from v_active_slide where file_id=?", fid)["c"] == 2


# ------------------------------------------------------------------ 矩阵：重复上传复用
def test_duplicate_upload_reuses(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "a.pptx", 2)
    f1 = upload(cfg, store, data, "a.pptx", "sample")
    f2 = upload(cfg, store, data, "b.pptx", "sample")
    assert f1 == f2
    assert db.one("select count(*) c from file")["c"] == 1


# ------------------------------------------------------------------ 矩阵：重跑 409
def test_active_conflict_409(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "a.pptx", 1)
    fid = upload(cfg, store, data)
    with pytest.raises(Exception):   # uq_ex_one_active
        with db.tx() as cur:
            cur.execute("insert into extraction(file_id,status) values(?, 'queued')", (fid,))


# ------------------------------------------------------------------ 矩阵：删除（终态守卫）
def test_delete_guard(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "a.pptx", 2)
    fid = upload(cfg, store, data)
    with pytest.raises(Conflict):    # queued 中删除 → 409
        store.delete_file(fid)
    w = Worker(db, cfg)
    drain(w)
    store.delete_file(fid)           # done 后可删
    assert db.one("select * from file where id=?", fid) is None
    assert db.one("select count(*) c from slide")["c"] == 0


# ------------------------------------------------------------------ 矩阵：running 崩溃 → 启动恢复
def test_startup_recovery(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "a.pptx", 2)
    fid = upload(cfg, store, data)
    ex = db.one("select id from extraction where file_id=?", fid)
    with db.tx() as cur:             # 模拟领取后崩溃
        cur.execute("update extraction set status='running' where id=?", (ex["id"],))
        cur.execute("insert into slide(extraction_id,file_id,no,png,thumb) values(?,?,1,'x','y')",
                    (ex["id"], fid))
    store.startup_recover_and_gc()
    ex2 = db.one("select * from extraction where id=?", ex["id"])
    assert ex2["status"] == "failed"
    assert ex2["error"] == "进程中断"
    assert db.one("select count(*) c from slide")["c"] == 0   # 残留清理


# ------------------------------------------------------------------ 矩阵：mark_failed 幂等 + 迟到回调
def test_mark_failed_idempotent(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "a.pptx", 1)
    fid = upload(cfg, store, data)
    ex = db.one("select id from extraction where file_id=?", fid)
    with db.tx() as cur:
        cur.execute("update extraction set status='running' where id=?", (ex["id"],))
    store.mark_failed(ex["id"], "boom")
    assert db.one("select status from extraction where id=?", ex["id"])["status"] == "failed"
    store.mark_failed(ex["id"], "late-callback")   # 幂等退出
    assert db.one("select error from extraction where id=?", ex["id"])["error"] == "boom"


# ------------------------------------------------------------------ 矩阵：finalize 幂等 / InvariantViolation
def test_finalize_idempotent(env):
    cfg, db, store = env
    data = make_pptx(cfg.tmp_dir / "a.pptx", 2)
    fid = upload(cfg, store, data)
    w = Worker(db, cfg)
    drain(w)
    ex = db.one("select * from extraction where file_id=?", fid)
    # 幂等：done 且（若 LO 可用）目录完整 → 正常返回；目录缺失 → InvariantViolation（不改状态）
    target = cfg.previews_dir / str(ex["id"])
    if target.exists():
        from ppt_lite.fslock import to_trash
        to_trash(target, cfg.trash_dir)          # 模拟目录丢失（与 LO 是否安装无关）
    with pytest.raises(InvariantViolation):
        store.finalize(ex["id"], 2)
    assert db.one("select status from extraction where id=?", ex["id"])["status"] == "done"  # 状态未被改动


# ------------------------------------------------------------------ 矩阵：媒体共享与 gc
def test_media_shared_across_files(env):
    """媒体引用计数 gc（不经 upload/worker——最小数据直入，聚焦 gc 语义）。"""
    from ppt_lite.pipeline import _stage_media
    cfg, db, store = env
    png = _png_bytes()
    m = _stage_media(cfg, cfg.tmp_dir / "mstage", png, ".png")
    mid = store.upsert_media(m["sha256"], m["fmt"], m["w"], m["h"], m["phash"], m["path"], None,
                             cfg.tmp_dir / "mstage" / Path(m["path"]).name)
    # 两个独立 file，各挂一条引用同一媒体
    fids: list[int] = []
    for i in range(2):
        with db.tx() as cur:
            cur.execute("insert into file(sha256,name,category,raw_path) values(?,?, 'sample',?)",
                        (f"sha-{i}", f"f{i}.pptx", f"raw/f{i}.pptx"))
            fid = cur.lastrowid
            cur.execute("insert into extraction(file_id,status) values(?, 'done')", (fid,))
            exid = cur.lastrowid
            cur.execute("insert into slide(extraction_id,file_id,no,png,thumb) values(?,?,1,'p','t')", (exid, fid))
            sid = cur.lastrowid
            cur.execute("insert into media_ref values(?,?, 'content')", (sid, mid))
            fids.append(fid)
    store.delete_file(fids[0])   # 删 f1：媒体仍被 f2 引用 → 保留
    assert db.one("select * from media where id=?", mid) is not None
    store.delete_file(fids[1])   # 删 f2：无引用 → media 行删除
    assert db.one("select * from media where id=?", mid) is None
    assert not (cfg.data_dir / m["path"]).exists()   # 规范文件已搬 trash


def _png_bytes() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


# ------------------------------------------------------------------ 预检拒收
def test_precheck_rejects():
    # 加密 CFB（OLE 容器，伪造 EncryptionInfo 太复杂——用 olefile 创建不现实；改为拒收未知魔数）
    with pytest.raises(PrecheckError):
        precheck(b"\x00\x01\x02garbage", soffice_available=False)
    with pytest.raises(PrecheckError):
        precheck(b"", soffice_available=False)


def test_tags_norm():
    from ppt_lite.db import norm_tags
    assert norm_tags("a, b ,,a，c ") == ",a,b,c,"
    assert norm_tags(None) == ","


# ------------------------------------------------------------------ 审核端点回归（防 upsert 括号笔误）
def test_review_endpoints(env):
    """/api/slides/{id}/review 与 /api/media/{sha}/review 的 upsert 语法回归。"""
    from fastapi.testclient import TestClient
    from ppt_lite.app import app
    cfg, db, store = env
    # 造素材行
    with db.tx() as cur:
        cur.execute("insert into media(sha256, fmt, path) values('m1','png','media/m1.png')")
    # 造文件+done+slide
    data = make_pptx(cfg.tmp_dir / "r.pptx", 1)
    fid = upload(cfg, store, data)
    Worker(db, cfg)  # noqa
    drain(Worker(db, cfg))
    sid = db.one("select id from v_active_slide where file_id=?", fid)["id"]

    c = TestClient(app)
    r = c.post(f"/api/slides/{sid}/review",
               data={"status": "approved", "slide_type": "cover", "tags": "战略,宣讲",
                     "quality": 5, "is_template": "true",
                     "template_json": '{"capacity": {"title_max_chars": 26}}'})
    assert r.status_code == 200, r.text
    # 断言用 app 的全局 db（端点写入方），与 env 的 db 在首个 import 时同源
    import ppt_lite.app as appmod
    adb = appmod.db
    sr = adb.one("select * from slide_review where file_id=?", fid)
    assert sr["status"] == "approved" and sr["is_template"] == 1 and sr["tags"] == ",战略,宣讲,"
    # upsert 二次更新
    r = c.post(f"/api/slides/{sid}/review", data={"status": "reference"})
    assert r.status_code == 200

    r = c.post("/api/media/m1/review", data={"status": "approved", "tags": "logo"})
    assert r.status_code == 200, r.text
    r = c.post("/api/media/m1/review", data={"status": "forbidden"})
    assert r.status_code == 200
    assert adb.one("select status from media_review where sha256='m1'")["status"] == "forbidden"


# ------------------------------------------------------------------ 超时自适应
def test_lo_timeout_scaling(env):
    cfg, _, _ = env
    base = cfg.lo_timeout(pages=0, size_bytes=0)                      # 基础开销
    assert base == cfg.lo_timeout_base_sec
    t_small = cfg.lo_timeout(pages=10, size_bytes=2 * 1024 * 1024)    # 10 页 2MB
    t_large = cfg.lo_timeout(pages=100, size_bytes=50 * 1024 * 1024)  # 100 页 50MB
    assert t_large > t_small > base
    assert cfg.lo_timeout(pages=10000, size_bytes=500 * 1024 * 1024) == cfg.lo_timeout_max_sec  # 封顶


# ------------------------------------------------------------------ 规范 markdown 流程
def test_design_markdown_flow(env):
    """规范承载改为 markdown：导入（粘贴）→ 编辑 → 确认 → 导出 .md。
    注意：app 模块级全局 db 在首次 import 时定型，断言必须用 appmod.db（与端点同源）。"""
    from fastapi.testclient import TestClient
    import ppt_lite.app as appmod
    c = TestClient(appmod.app)
    adb = appmod.db

    # 粘贴导入
    md = "# 品牌规范\n\n## 配色\n- 主色 #003A70\n- 禁用 #FF0000\n\n## 字体\n- 标题 MiSans 26pt 加粗\n"
    r = c.post("/api/design/import", data={"content": md})
    assert r.status_code == 200, r.text
    v = r.json()["version"]
    row = adb.one("select * from design_system where version=?", v)
    assert row["status"] == "draft" and row["content_md"] == md.strip()

    # 上传 .md 文件导入
    r = c.post("/api/design/import", files={"file": ("spec.md", "# v2 规范\n字体规范…", "text/markdown")})
    assert r.status_code == 200 and r.json()["version"] == v + 1

    # 材料类型（file.doc_type）：上传携带 + 后续可改
    import tempfile, pathlib
    r2 = c.post("/api/upload",
                files={"file": ("x.pptx", make_pptx(pathlib.Path(tempfile.mkdtemp()) / "dt.pptx", 1), "application/octet-stream")},
                data={"category": "sample", "doc_type": "洞察材料"})
    assert r2.status_code == 200, r2.text
    fid = r2.json()["file_id"]
    assert adb.one("select doc_type from file where id=?", fid)["doc_type"] == "洞察材料"
    r3 = c.post(f"/api/files/{fid}/attrs", data={"doc_type": "BP材料"})
    assert r3.status_code == 200
    assert adb.one("select doc_type from file where id=?", fid)["doc_type"] == "BP材料"

    # 编辑草稿（markdown 路径）
    did = row["id"]
    r = c.post(f"/api/design/{did}/update", data={"content": md + "\n## 补充\n- 页脚必须\n"})
    assert r.status_code == 200
    assert "页脚必须" in adb.one("select content_md from design_system where id=?", did)["content_md"]

    # 确认后冻结
    r = c.post(f"/api/design/{did}/confirm")
    assert r.status_code == 200
    r = c.post(f"/api/design/{did}/update", data={"content": "改不动的"})
    assert r.status_code == 409

    # 导出 .md
    r = c.get(f"/api/design/{did}/export")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/markdown")
    assert "design-spec-v" in r.headers["content-disposition"]
