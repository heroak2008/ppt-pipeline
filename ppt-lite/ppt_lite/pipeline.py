"""pipeline.py：解析+渲染 → ExtractedDocument（§7 契约；解析器不直接写库）。"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import imagehash
from PIL import Image

from . import PIPELINE_VERSION
from .config import Config
from .render import convert_to_pdf, render_pdf_pages

log = logging.getLogger("pptlite.pipeline")

# 媒体扩展名白名单（zip entry 原样保留的格式）
MEDIA_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".svg",
              ".emf", ".wmf", ".webp", ".ico"}
# 需要 PNG 预览的矢量/未支持格式
PREVIEW_EXTS = {".emf", ".wmf", ".svg"}


@dataclass
class ExtractedDocument:
    """契约 v1（§7）。slides[].png/thumb 为最终路径 previews/{ex_id}/...（发布后生效）。"""
    source: dict
    versions: dict = field(default_factory=lambda: {"parser": PIPELINE_VERSION})
    slides: list[dict] = field(default_factory=list)
    media: list[dict] = field(default_factory=list)
    occurrences: list[dict] = field(default_factory=list)
    design: dict = field(default_factory=dict)
    render_error: str | None = None
    warnings: list[str] = field(default_factory=list)

    def to_json(self) -> str:
        return json.dumps(self.__dict__, ensure_ascii=False)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_document(cfg: Config, ex_id: int, src_path: Path, category: str,
                     name: str, raw_rel: str, derived_rel: str | None) -> ExtractedDocument:
    """输入文件 + 工作目录，产出契约 dict。全部锁外。"""
    doc = ExtractedDocument(
        source={"name": name, "sha256": "", "category": category,
                "raw_path": raw_rel, "derived_path": derived_rel},
    )
    doc.source["sha256"] = sha256_bytes(src_path.read_bytes())

    work = cfg.tmp_dir / str(ex_id)
    work.mkdir(parents=True, exist_ok=True)
    (work / "pdf").mkdir(exist_ok=True)
    (work / "media_staging").mkdir(exist_ok=True)

    # ── unpacking：媒体原样落 media_staging（sha256 命名）+ 主题 ──
    media_by_part: dict[str, dict] = {}   # part name -> media dict
    try:
        with zipfile.ZipFile(src_path) as zf:
            total = 0
            for info in zf.infolist():
                total += info.file_size
                if total > cfg.max_unzip_total_bytes:
                    doc.warnings.append(f"解压总量超限（>{cfg.max_unzip_total_bytes}B），跳过剩余媒体")
                    break
                name_l = info.filename.lower()
                ext = Path(name_l).suffix
                if name_l.startswith("ppt/media/") and ext in MEDIA_EXTS and info.file_size <= cfg.max_media_bytes:
                    data = zf.read(info)
                    m = _stage_media(cfg, work / "media_staging", data, ext)
                    media_by_part[info.filename] = m
                    if not any(x["sha256"] == m["sha256"] for x in doc.media):
                        doc.media.append(m)
            # 主题色/字体（spec 草稿数据源）
            theme = _read_theme(zf)
            if theme:
                doc.design["theme_colors"] = theme.get("colors", [])
                doc.design["theme_fonts"] = theme.get("fonts", [])
    except zipfile.BadZipFile:
        doc.warnings.append("ZIP 打开失败（应已在预检拦截）")

    # ── parsing：python-pptx 逐页 ──
    from pptx import Presentation
    from pptx.util import Emu

    try:
        prs = Presentation(str(src_path))
        doc.design["page_size"] = {"width_emu": prs.slide_width, "height_emu": prs.slide_height}
        slide_media_map: dict[int, list[tuple[str, str]]] = {}   # slide_no -> [(part, role)]
        font_counter: Counter = Counter()                        # (字体, 字号pt, bold) -> 次数

        for idx, slide in enumerate(prs.slides, start=1):
            s: dict = {"no": idx, "title": None, "texts": [], "struct": {}, "notes": None,
                       "hidden": False, "warnings": [], "png": f"previews/{ex_id}/slides/{idx}.png",
                       "thumb": f"previews/{ex_id}/thumbs/{idx}.png"}
            texts_parts: list[str] = []
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                    s["notes"] = slide.notes_slide.notes_text_frame.text
                for shp in slide.shapes:
                    _walk_shape(shp, s, texts_parts, slide_media_map, idx, font_counter)
            except Exception as e:  # noqa: BLE001 — 页级失败记 warnings，不中断（§4.3）
                s["warnings"].append(f"页解析异常: {e!r}")
            s["title"] = s["title"] or (texts_parts[0][:80] if texts_parts else None)
            s["search_text"] = (s["title"] or "") + "\n" + "\n".join(texts_parts)
            doc.slides.append(s)

        for no, refs in slide_media_map.items():
            for part, role in refs:
                m = media_by_part.get(part)
                if m:
                    doc.occurrences.append({"slide_no": no, "media_sha256": m["sha256"], "role": role})

        # ── 实采字体统计（run 级）与形状填充色直方图 → 规范草稿数据源 ──
        doc.design["font_usage"] = [
            {"font": f, "size_pt": sz, "bold": b, "count": n}
            for (f, sz, b), n in font_counter.most_common(30) if f or sz
        ]
        color_counter: Counter = Counter()
        for sl in doc.slides:
            for e in (sl.get("struct", {}).get("shapes") or []):
                fill = e.get("fill")
                if fill and fill.startswith("#"):
                    color_counter[fill] += 1
        doc.design["fill_usage"] = [{"rgb": c, "count": n} for c, n in color_counter.most_common(20)]
    except Exception as e:  # noqa: BLE001 — 文档级解析失败
        doc.warnings.append(f"python-pptx 解析失败: {e!r}")

    # ── rendering：LO→PDF→PNG（LO 缺失/失败 → render_error 降级，png 占位仍写 DB）──
    if doc.slides:
        try:
            if not cfg.soffice:
                doc.render_error = "LibreOffice 未配置，跳过渲染（预览不可用，数据完整）"
            else:
                # 超时按页数 + 文件大小动态计算（页数在解析后已知）
                timeout = cfg.lo_timeout(pages=len(doc.slides), size_bytes=src_path.stat().st_size)
                log.info("LO 渲染超时窗口: %ss（%d 页 / %.1fMB）",
                         timeout, len(doc.slides), src_path.stat().st_size / 1048576)
                pdf, render_reason = convert_to_pdf(src_path, work / "pdf", cfg.soffice, timeout)
                if pdf is None:
                    doc.render_error = render_reason or "LO 转换失败（原因未知）"
                else:
                    render_pdf_pages(pdf, work / "preview" / "slides", work / "preview" / "thumbs",
                                     cfg.slide_png_width, cfg.thumb_width)
        except Exception as e:  # noqa: BLE001
            doc.render_error = f"渲染异常: {e!r}"

    return doc


def _stage_media(cfg: Config, staging: Path, data: bytes, ext: str) -> dict:
    """媒体暂存落盘（sha256 命名）+ 基本信息（w/h/phash）。"""
    staging.mkdir(parents=True, exist_ok=True)
    sha = sha256_bytes(data)
    fname = f"{sha}{ext}"
    fpath = staging / fname
    if not fpath.exists():
        fpath.write_bytes(data)
    m: dict = {"sha256": sha, "fmt": ext.lstrip(".").lower(), "w": None, "h": None,
               "path": f"media/{fname}", "preview": None, "phash": None}
    try:
        if ext.lower() in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff", ".ico"}:
            img = Image.open(io.BytesIO(data))
            m["w"], m["h"] = img.size
            m["phash"] = str(imagehash.phash(img))
    except Exception:  # noqa: BLE001
        pass
    return m


def _walk_shape(shp, slide_dict: dict, texts_parts: list, slide_media_map: dict, slide_no: int,
                font_counter: Counter | None = None) -> None:
    """遍历形状（含组合递归）：全量几何采集 + 文本/占位符/媒体引用；SmartArt/OLE 记 presence。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    # ── 全量形状采集（矢量版式的数据基础：色块/线条/组合的几何+填充+描边）──
    shapes = slide_dict["struct"].setdefault("shapes", [])
    if len(shapes) < 300:                     # 上限防超大页膨胀
        shapes.append(_shape_entry(shp))

    st = shp.shape_type
    try:
        if st == MSO_SHAPE_TYPE.GROUP:
            for sub in shp.shapes:
                _walk_shape(sub, slide_dict, texts_parts, slide_media_map, slide_no, font_counter)
            return
    except Exception:  # noqa: BLE001
        pass

    if shp.has_text_frame:
        txt = shp.text_frame.text.strip()
        if txt:
            texts_parts.append(txt)
            slide_dict["texts"].append({
                "name": shp.name, "type": "textbox",
                "text": txt,
                "left": shp.left, "top": shp.top, "width": shp.width, "height": shp.height,
            })
            if shp.is_placeholder and slide_dict["title"] is None and _is_title_placeholder(shp):
                slide_dict["title"] = txt
        if font_counter is not None:
            _collect_fonts(shp.text_frame, font_counter)
    try:
        if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = shp.image
            part = image.filename or ""
            blob_sha = sha256_bytes(image.blob)
            role = "content"
            slide_media_map.setdefault(slide_no, []).append((_find_part_by_sha(shp, blob_sha), role))
    except Exception:  # noqa: BLE001
        pass
    # SmartArt/OLE/图表 presence（不拆解）
    xml = getattr(shp, "_element", None)
    if xml is not None:
        tag = xml.tag
        if "graphicFrame" in tag:
            body = "".join(e.tag for e in xml.iter())
            if "diagram" in body:
                slide_dict["struct"].setdefault("has_smartart", 0)
                slide_dict["struct"]["has_smartart"] += 1
            elif "chart" in body:
                slide_dict["struct"].setdefault("has_chart", 0)
                slide_dict["struct"]["has_chart"] += 1
            elif "oleObject" in body:
                slide_dict["struct"].setdefault("has_ole", 0)
                slide_dict["struct"]["has_ole"] += 1
    if getattr(shp, "has_table", False):
        try:
            tbl = shp.table
            cells = [[c.text for c in row.cells] for row in tbl.rows]
            slide_dict["struct"].setdefault("tables", []).append(cells)
            texts_parts.extend(t for row in cells for t in row if t)
        except Exception:  # noqa: BLE001
            pass


def _collect_fonts(text_frame, font_counter: Counter) -> None:
    """run 级字体统计：字体名 / 字号 pt / 加粗（继承主题字体的 run 记为 None 字段，由草稿层合并解释）。"""
    try:
        for para in text_frame.paragraphs:
            for run in para.runs:
                fname = run.font.name
                size = round(run.font.size.pt, 1) if run.font.size else None
                bold = bool(run.font.bold) if run.font.bold is not None else None
                if fname or size:
                    font_counter[(fname, size, bold)] += 1
    except Exception:  # noqa: BLE001
        pass


def _shape_entry(shp) -> dict:
    """单个形状的几何+样式记录（EMU 坐标；颜色尽量解析为 RGB）。"""
    e: dict = {
        "type": str(shp.shape_type).split(" ")[0] if shp.shape_type is not None else "UNKNOWN",
        "name": shp.name,
        "x": shp.left, "y": shp.top, "w": shp.width, "h": shp.height,
    }
    try:
        if shp.rotation:
            e["rot"] = shp.rotation
    except Exception:  # noqa: BLE001
        pass
    for key, attr in (("fill", "fill"), ("line", "line")):
        try:
            fmt = getattr(shp, attr, None)
            if fmt is None:
                continue
            ftype = fmt.type
            if ftype is not None and str(ftype).startswith("SOLID"):
                color = fmt.fore_color
                if color is not None and "RGB" in str(getattr(color, "type", "")):
                    e[key] = f"#{color.rgb}"
                else:
                    # 主题色槽位（如 ACCENT_1）——保留槽位名，主题色映射由 theme 提供
                    try:
                        e[key] = f"theme:{color.theme_color}"
                    except Exception:  # noqa: BLE001
                        pass
        except Exception:  # noqa: BLE001
            pass
    return e


def _is_title_placeholder(shp) -> bool:
    from pptx.enum.shapes import PP_PLACEHOLDER
    try:
        return shp.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    except Exception:  # noqa: BLE001
        return False


def _find_part_by_sha(shp, blob_sha: str) -> str:
    """由形状的 rId 反查媒体 part 名（找不到则回退 sha 定位键）。"""
    try:
        rid = shp._element.blip_rId
        part = shp.part.related_part(rid)
        return part.partname
    except Exception:  # noqa: BLE001
        return f"sha:{blob_sha}"


def _read_theme(zf: zipfile.ZipFile) -> dict | None:
    """读 theme1.xml：主题色（槽位→RGB）与主题字体。"""
    import lxml.etree as ET

    names = [n for n in zf.namelist() if n.startswith("ppt/theme/") and n.endswith(".xml")]
    if not names:
        return None
    try:
        root = ET.fromstring(zf.read(names[0]))
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        colors = []
        for el in root.findall(".//a:clrScheme/*", ns):
            tag = el.tag.split("}")[1]
            srgb = el.find("a:srgbClr", ns)
            sysc = el.find("a:sysClr", ns)
            val = srgb.get("val") if srgb is not None else (sysc.get("lastClr") if sysc is not None else None)
            if val:
                colors.append({"slot": tag, "rgb": f"#{val}"})
        fonts = []
        for el in root.findall(".//a:fontScheme/*/a:latin", ns):
            fonts.append({"typeface": el.get("typeface")})
        major = root.find(".//a:majorFont/a:latin", ns)
        minor = root.find(".//a:minorFont/a:latin", ns)
        return {"colors": colors,
                "fonts": [{"major": major.get("typeface") if major is not None else None,
                           "minor": minor.get("typeface") if minor is not None else None}]}
    except Exception:  # noqa: BLE001
        return None
